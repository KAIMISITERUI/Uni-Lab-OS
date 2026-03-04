#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
功能:
    验证多工作表场景下的任务解析与产率配置解析.
参数:
    无.
返回:
    无.
"""

import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import openpyxl

from eit_analysis_station.config.setting import Settings
from eit_analysis_station.controller.analysis_controller import AnalysisStationController
from eit_analysis_station.processor.yield_calculator import (
    YieldCalculator,
    YIELD_CONFIG_SHEET_NAME,
)


class TestMultiSheetParsing(unittest.TestCase):
    """
    功能:
        覆盖分析站在多工作表下的关键解析路径.
    参数:
        无.
    返回:
        无.
    """

    @staticmethod
    def _make_tmp_dir() -> Path:
        """
        功能:
            创建测试临时目录.
        参数:
            无.
        返回:
            Path, 临时目录路径.
        """
        tmp_root = Path.cwd() / "eit_analysis_station" / "tests" / "_tmp"
        tmp_root.mkdir(parents=True, exist_ok=True)
        case_dir = tmp_root / f"multisheet_{uuid4().hex}"
        case_dir.mkdir(parents=True, exist_ok=False)
        return case_dir

    def test_parse_task_xlsx_prefers_plan_sheet_when_active_is_yield_sheet(self) -> None:
        """
        功能:
            验证 active 指向非任务配置页时, 仍能解析到实验编号和 GC_MS 方法.
        参数:
            无.
        返回:
            无.
        """
        case_dir = self._make_tmp_dir()
        task_dir = case_dir / "tasks" / "754"
        task_dir.mkdir(parents=True, exist_ok=True)
        plan_path = task_dir / "754_experiment_plan.xlsx"

        workbook = openpyxl.Workbook()
        ws_plan = workbook.active
        ws_plan.title = "实验方案设定"
        ws_plan.cell(row=1, column=1, value="实验设定")
        ws_plan.cell(row=1, column=3, value="实验编号")
        for row in range(2, 14):
            ws_plan.cell(row=row, column=3, value=row - 1)
        ws_plan.cell(row=31, column=1, value="GC_MS")
        ws_plan.cell(row=31, column=2, value="test")
        ws_plan.cell(row=32, column=1, value="UPLC_QTOF")
        ws_plan.cell(row=33, column=1, value="HPLC")

        ws_yield = workbook.create_sheet("GC产率计算")
        ws_yield.cell(row=8, column=1, value="适用实验")
        ws_yield.cell(row=9, column=3, value=1)

        workbook.active = 1  # active 指向非任务配置页
        workbook.save(plan_path)
        workbook.close()

        controller = AnalysisStationController(
            settings=Settings(synthesis_tasks_dir=case_dir / "tasks")
        )
        parsed = controller._parse_task_xlsx(task_dir, "754")

        self.assertEqual(parsed["exp_count"], 12)
        self.assertEqual(parsed["gc_ms_method"], "test")
        self.assertIsNone(parsed["uplc_qtof_method"])
        self.assertIsNone(parsed["hplc_method"])

    def test_parse_yield_config_reads_main_params_from_non_first_sheet(self) -> None:
        """
        功能:
            验证主参数不在第一张工作表时, parse_yield_config 仍可正确读取.
        参数:
            无.
        返回:
            无.
        """
        case_dir = self._make_tmp_dir()
        plan_path = case_dir / "plan.xlsx"
        chemical_path = case_dir / "chemical_list.xlsx"

        workbook = openpyxl.Workbook()
        ws_misc = workbook.active
        ws_misc.title = "其它说明"
        ws_misc.cell(row=1, column=1, value="说明")

        ws_main = workbook.create_sheet("实验方案设定")
        ws_main.cell(row=1, column=1, value="反应规模(mmol)")
        ws_main.cell(row=1, column=2, value=0.2)
        ws_main.cell(row=2, column=1, value="内标种类")
        ws_main.cell(row=2, column=2, value="IS")
        ws_main.cell(row=3, column=1, value="内标用量(μL/mg)")
        ws_main.cell(row=3, column=2, value=50)

        ws_yield = workbook.create_sheet(YIELD_CONFIG_SHEET_NAME)
        ws_yield.cell(row=1, column=1, value="内标SMILES")
        ws_yield.cell(row=1, column=2, value="CC")
        ws_yield.cell(row=2, column=1, value="内标预期RT(min)")
        ws_yield.cell(row=2, column=2, value=6.8)
        ws_yield.cell(row=3, column=1, value="产率计算方法")
        ws_yield.cell(row=3, column=2, value="ECN")
        ws_yield.cell(row=5, column=1, value="适用实验")
        ws_yield.cell(row=5, column=2, value="目标产物名称")
        ws_yield.cell(row=5, column=3, value="SMILES")
        ws_yield.cell(row=5, column=4, value="当量(eq)")
        ws_yield.cell(row=6, column=1, value="1-12")
        ws_yield.cell(row=6, column=2, value="产物A")
        ws_yield.cell(row=6, column=3, value="CC")
        ws_yield.cell(row=6, column=4, value=1)

        workbook.active = 0
        workbook.save(plan_path)
        workbook.close()

        chem_workbook = openpyxl.Workbook()
        chem_sheet = chem_workbook.active
        chem_sheet.title = "chemical_list"
        chem_sheet.append(
            [
                "substance",
                "molecular_weight",
                "physical_state",
                "density (g/mL)",
                "physical_form",
                "active_content",
                "storage_location",
            ]
        )
        chem_sheet.append(["IS", 100, "liquid", 1.0, "solution", "1 mmol/mL", "L1"])
        chem_workbook.save(chemical_path)
        chem_workbook.close()

        calculator = YieldCalculator()
        with patch.object(calculator, "_calculate_is_moles", return_value=1e-4):
            config = calculator.parse_yield_config(plan_path, chemical_path)

        self.assertAlmostEqual(config.reaction_scale_mmol, 0.2, places=6)
        self.assertEqual(config.is_name, "IS")


if __name__ == "__main__":
    unittest.main()

