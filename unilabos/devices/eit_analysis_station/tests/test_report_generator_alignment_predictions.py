#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
功能:
    验证 TIC-FID 对照表新增预测列与名称超链接写入行为.
参数:
    无.
返回:
    无.
"""

import unittest
from pathlib import Path
from uuid import uuid4

import openpyxl

from eit_analysis_station.processor.nist_matcher import CompoundMatch
from eit_analysis_station.processor.peak_integrator import PeakResult
from eit_analysis_station.processor.report_generator import (
    PIMPrediction,
    SSHMPrediction,
    iHSHMPrediction,
    ReportGenerator,
    SampleResult,
)


class TestReportGeneratorAlignmentPredictions(unittest.TestCase):
    """
    功能:
        覆盖 TIC-FID 对照表的列头、预测结果和名称超链接写入.
    参数:
        无.
    返回:
        无.
    """

    @staticmethod
    def _make_tmp_dir() -> Path:
        """
        功能:
            在仓库可写目录创建临时测试目录.
        参数:
            无.
        返回:
            Path, 临时目录路径.
        """
        tmp_root = Path.cwd() / "eit_analysis_station" / "tests" / "_tmp"
        tmp_root.mkdir(parents=True, exist_ok=True)
        case_dir = tmp_root / f"alignment_{uuid4().hex}"
        case_dir.mkdir(parents=True, exist_ok=False)
        return case_dir

    @staticmethod
    def _build_peak(peak_index: int, retention_time: float, area: float) -> PeakResult:
        """
        功能:
            构造最小峰对象.
        参数:
            peak_index: 峰号.
            retention_time: 保留时间(min).
            area: 峰面积.
        返回:
            PeakResult.
        """
        return PeakResult(
            peak_index=peak_index,
            retention_time=retention_time,
            height=1000.0,
            area=area,
            area_percent=100.0,
            start_time=retention_time - 0.1,
            end_time=retention_time + 0.1,
            width=0.2,
        )

    @classmethod
    def _build_sample_ok(cls) -> SampleResult:
        """
        功能:
            构建包含预测结果与 Top2 命中的样品对象.
        参数:
            无.
        返回:
            SampleResult.
        """
        tic_peak = cls._build_peak(1, 5.000, 5000.0)
        fid_peak = cls._build_peak(1, 5.002, 3000.0)
        match_1 = CompoundMatch(
            compound_name="Benzylamine",
            cas_number="",
            match_score=92.0,
            formula="C7H9N",
            mw=107.0,
            nist_id=22326,
        )
        match_2 = CompoundMatch(
            compound_name="Aniline, N-methyl-",
            cas_number="100-61-8",
            match_score=85.0,
            formula="C7H9N",
            mw=107.0,
        )
        return SampleResult(
            sample_name="sample-1",
            d_dir=Path("sample-1.D"),
            tic_peaks=[tic_peak],
            fid_peaks=[fid_peak],
            compound_matches={5.000: [match_1, match_2]},
            pim_predictions={
                5.001: PIMPrediction(
                    predicted_mw=180,
                    confidence_index=1.23456,
                    status="ok",
                    message="PIM 预测成功",
                )
            },
            sshm_predictions={
                5.001: SSHMPrediction(
                    predicted_mw=182,
                    confidence=0.8765,
                    correction=2,
                    status="ok",
                    message="SS-HM 预测成功",
                )
            },
            ihshm_predictions={
                5.001: iHSHMPrediction(
                    predicted_mw=183,
                    confidence=0.012345,
                    status="ok",
                    message="iHS-HM 预测成功",
                )
            },
        )

    @classmethod
    def _build_sample_fid_only(cls) -> SampleResult:
        """
        功能:
            构建仅有 FID 峰的样品对象.
        参数:
            无.
        返回:
            SampleResult.
        """
        return SampleResult(
            sample_name="sample-fid-only",
            d_dir=Path("sample-fid-only.D"),
            fid_peaks=[cls._build_peak(1, 6.200, 1200.0)],
        )

    @classmethod
    def _build_sample_error_prediction(cls) -> SampleResult:
        """
        功能:
            构建预测状态为 error 的样品对象.
        参数:
            无.
        返回:
            SampleResult.
        """
        tic_peak = cls._build_peak(1, 7.000, 4200.0)
        fid_peak = cls._build_peak(1, 7.003, 1800.0)
        return SampleResult(
            sample_name="sample-error",
            d_dir=Path("sample-error.D"),
            tic_peaks=[tic_peak],
            fid_peaks=[fid_peak],
            pim_predictions={
                7.001: PIMPrediction(
                    predicted_mw=None,
                    confidence_index=None,
                    status="error",
                    message="PIM 失败",
                )
            },
            sshm_predictions={
                7.001: SSHMPrediction(
                    predicted_mw=None,
                    confidence=None,
                    status="error",
                    message="SS-HM 失败",
                )
            },
            ihshm_predictions={
                7.001: iHSHMPrediction(
                    predicted_mw=None,
                    confidence=None,
                    status="error",
                    message="iHS-HM 失败",
                )
            },
        )

    def test_alignment_sheet_headers_values_and_hyperlinks(self) -> None:
        """
        功能:
            验证新增列列头、预测值、结构图和质谱图超链接写入.
        参数:
            无.
        返回:
            无.
        """
        sample = self._build_sample_ok()
        tmp_dir = self._make_tmp_dir()
        nist_png = tmp_dir / "NIST_22326.png"
        cas_png = tmp_dir / "CAS_100618.png"
        ms_png = tmp_dir / "sample_1_ms.png"
        nist_png.write_bytes(b"png")
        cas_png.write_bytes(b"png")
        ms_png.write_bytes(b"png")
        sample.ms_plot_paths = {1: ms_png}

        workbook = openpyxl.Workbook()
        ws = workbook.active

        ReportGenerator()._write_alignment_sheet(
            ws,
            [sample],
            structure_images={
                "NIST:22326": nist_png,
                "CAS:100618": cas_png,
            },
        )

        self.assertEqual(ws.cell(row=1, column=15).value, "PIM预测分子量(Da)")
        self.assertEqual(ws.cell(row=1, column=16).value, "PIM置信指数")
        self.assertEqual(ws.cell(row=1, column=17).value, "SS-HM预测分子量(Da)")
        self.assertEqual(ws.cell(row=1, column=18).value, "SS-HM置信度")
        self.assertEqual(ws.cell(row=1, column=19).value, "iHS-HM预测分子量(Da)")
        self.assertEqual(ws.cell(row=1, column=20).value, "iHS-HM置信度")
        self.assertEqual(ws.cell(row=1, column=21).value, "质谱图")

        name_cell_1 = ws.cell(row=2, column=7)
        name_cell_2 = ws.cell(row=2, column=11)
        self.assertEqual(name_cell_1.value, "Benzylamine")
        self.assertEqual(name_cell_2.value, "Aniline, N-methyl-")
        self.assertEqual(name_cell_1.hyperlink.target, str(nist_png))
        self.assertEqual(name_cell_2.hyperlink.target, str(cas_png))

        self.assertEqual(ws.cell(row=2, column=15).value, 180)
        self.assertAlmostEqual(ws.cell(row=2, column=16).value, 1.2346, places=4)
        self.assertEqual(ws.cell(row=2, column=17).value, 182)
        self.assertAlmostEqual(ws.cell(row=2, column=18).value, 0.8765, places=4)
        self.assertEqual(ws.cell(row=2, column=19).value, 183)
        self.assertAlmostEqual(ws.cell(row=2, column=20).value, 0.012345, places=6)
        ms_cell = ws.cell(row=2, column=21)
        self.assertEqual(ms_cell.value, "查看质谱图")
        self.assertEqual(ms_cell.hyperlink.target, str(ms_png))
        workbook.close()

    def test_alignment_sheet_fid_only_row_keeps_prediction_columns_empty(self) -> None:
        """
        功能:
            验证仅 FID 峰行新增预测列保持空值.
        参数:
            无.
        返回:
            无.
        """
        sample = self._build_sample_fid_only()
        workbook = openpyxl.Workbook()
        ws = workbook.active

        ReportGenerator()._write_alignment_sheet(ws, [sample])

        for col in range(15, 22):
            self.assertIsNone(ws.cell(row=2, column=col).value)
        workbook.close()

    def test_alignment_sheet_error_predictions_keep_columns_empty(self) -> None:
        """
        功能:
            验证预测为 error 且无数值时, 新增预测列保持空值.
        参数:
            无.
        返回:
            无.
        """
        sample = self._build_sample_error_prediction()
        workbook = openpyxl.Workbook()
        ws = workbook.active

        ReportGenerator()._write_alignment_sheet(ws, [sample])

        for col in range(15, 22):
            self.assertIsNone(ws.cell(row=2, column=col).value)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
