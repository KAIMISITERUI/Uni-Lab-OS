#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
功能:
    验证 TIC 峰表中的 PIM/SS-HM/iHS-HM 预测列写入行为.
参数:
    无.
返回:
    无.
"""

import unittest
from pathlib import Path

import openpyxl

from eit_analysis_station.processor.peak_integrator import PeakResult
from eit_analysis_station.processor.report_generator import (
    PIMPrediction,
    SSHMPrediction,
    iHSHMPrediction,
    ReportGenerator,
    SampleResult,
)


class TestReportGeneratorPredictionColumns(unittest.TestCase):
    """
    功能:
        覆盖 PIM/SS-HM/iHS-HM 列头与结果写入.
    参数:
        无.
    返回:
        无.
    """

    @staticmethod
    def _build_sample_ok() -> SampleResult:
        """
        功能:
            构建包含 ok 状态预测的样品对象.
        参数:
            无.
        返回:
            SampleResult.
        """
        return SampleResult(
            sample_name="sample-1",
            d_dir=Path("sample-1.D"),
            tic_peaks=[
                PeakResult(
                    peak_index=1,
                    retention_time=5.000,
                    height=1000.0,
                    area=5000.0,
                    area_percent=100.0,
                    start_time=4.900,
                    end_time=5.100,
                    width=0.200,
                )
            ],
            pim_predictions={
                5.001: PIMPrediction(
                    predicted_mz=180,
                    predicted_mw=180,
                    confidence_index=1.23456,
                    status="ok",
                    message="PIM 预测成功",
                )
            },
            sshm_predictions={
                5.001: SSHMPrediction(
                    predicted_mw=182,
                    confidence=0.8765,
                    correction=2,
                    status="ok",
                    message="SS-HM 预测成功, 25 个命中",
                )
            },
            ihshm_predictions={
                5.001: iHSHMPrediction(
                    predicted_mw=183,
                    confidence=0.012345,
                    status="ok",
                    message="iHS-HM 预测成功, 范围 162-250 Da",
                )
            },
        )

    @staticmethod
    def _build_sample_error() -> SampleResult:
        """
        功能:
            构建包含 error 状态预测的样品对象.
        参数:
            无.
        返回:
            SampleResult.
        """
        return SampleResult(
            sample_name="sample-2",
            d_dir=Path("sample-2.D"),
            tic_peaks=[
                PeakResult(
                    peak_index=1,
                    retention_time=6.200,
                    height=800.0,
                    area=3200.0,
                    area_percent=100.0,
                    start_time=6.100,
                    end_time=6.300,
                    width=0.200,
                )
            ],
            pim_predictions={
                6.210: PIMPrediction(
                    predicted_mz=None,
                    predicted_mw=None,
                    confidence_index=None,
                    status="error",
                    message="PIM 置信指数计算失败",
                )
            },
            sshm_predictions={
                6.210: SSHMPrediction(
                    status="error",
                    message="MSPepSearch 不可用",
                )
            },
            ihshm_predictions={
                6.210: iHSHMPrediction(
                    status="error",
                    message="MSPepSearch 不可用",
                )
            },
        )

    def test_tic_sheet_contains_prediction_headers(self) -> None:
        """
        功能:
            验证 PIM/SS-HM/iHS-HM 列头存在且位置正确.
        """
        workbook = openpyxl.Workbook()
        ws = workbook.active

        ReportGenerator()._write_tic_sheet(ws, [self._build_sample_ok()])

        # PIM 列 (col 19-20)
        self.assertEqual(ws.cell(row=1, column=19).value, "PIM预测分子量(Da)")
        self.assertEqual(ws.cell(row=1, column=20).value, "PIM置信指数")

        # SS-HM 列 (col 21-22)
        self.assertEqual(ws.cell(row=1, column=21).value, "SS-HM预测分子量(Da)")
        self.assertEqual(ws.cell(row=1, column=22).value, "SS-HM置信度")

        # iHS-HM 列 (col 23-24)
        self.assertEqual(ws.cell(row=1, column=23).value, "iHS-HM预测分子量(Da)")
        self.assertEqual(ws.cell(row=1, column=24).value, "iHS-HM置信度")

        workbook.close()

    def test_tic_sheet_ok_values(self) -> None:
        """
        功能:
            验证 ok 状态下所有预测列写入正确数值.
        """
        workbook = openpyxl.Workbook()
        ws = workbook.active

        ReportGenerator()._write_tic_sheet(ws, [self._build_sample_ok()])

        # PIM: col 19=MW, col 20=confidence
        self.assertEqual(ws.cell(row=2, column=19).value, 180)
        self.assertAlmostEqual(ws.cell(row=2, column=20).value, 1.2346, places=4)

        # SS-HM: col 21=MW, col 22=confidence
        self.assertEqual(ws.cell(row=2, column=21).value, 182)
        self.assertAlmostEqual(ws.cell(row=2, column=22).value, 0.8765, places=4)

        # iHS-HM: col 23=MW, col 24=confidence
        self.assertEqual(ws.cell(row=2, column=23).value, 183)
        self.assertAlmostEqual(ws.cell(row=2, column=24).value, 0.012345, places=6)

        workbook.close()

    def test_tic_sheet_error_status_without_numeric_values(self) -> None:
        """
        功能:
            验证 error 状态下预测列无数值输出.
        """
        workbook = openpyxl.Workbook()
        ws = workbook.active

        ReportGenerator()._write_tic_sheet(ws, [self._build_sample_error()])

        # PIM: error 状态不写入数值
        self.assertIsNone(ws.cell(row=2, column=19).value)
        self.assertIsNone(ws.cell(row=2, column=20).value)

        # SS-HM: error 状态不写入数值
        self.assertIsNone(ws.cell(row=2, column=21).value)
        self.assertIsNone(ws.cell(row=2, column=22).value)

        # iHS-HM: error 状态不写入数值
        self.assertIsNone(ws.cell(row=2, column=23).value)
        self.assertIsNone(ws.cell(row=2, column=24).value)

        workbook.close()


if __name__ == "__main__":
    unittest.main()
