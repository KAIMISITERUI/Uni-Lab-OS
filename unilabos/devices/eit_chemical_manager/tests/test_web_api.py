# -*- coding: utf-8 -*-
"""
功能:
    覆盖 eit_chemical_manager Web API 的端到端行为.
    使用 FastAPI TestClient + tmp_path 提供的隔离 SQLite 数据库.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Dict

import openpyxl
import pytest
from fastapi.testclient import TestClient

from unilabos.devices.eit_chemical_manager.config.setting import Settings
from unilabos.devices.eit_chemical_manager.manager.chemical_manager import (
    ChemicalManager,
)
from unilabos.devices.eit_chemical_manager.web.app import create_app
from unilabos.devices.eit_chemical_manager.web.deps import get_manager


@pytest.fixture()
def manager(tmp_path: Path) -> ChemicalManager:
    """
    功能:
        创建隔离的 ChemicalManager 实例, 测试结束关闭数据库连接.
    """
    db_path = tmp_path / "chemical_library.db"
    data_dir = tmp_path / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    settings = Settings(db_path=db_path, data_dir=data_dir)
    inst = ChemicalManager(settings=settings)
    yield inst
    inst.db.close()


@pytest.fixture()
def client(manager: ChemicalManager, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    """
    功能:
        构造 TestClient 并把 get_manager 依赖替换为隔离实例.
        清空 CHEM_MGR_WEB_TOKEN 以便默认开放访问.
    """
    monkeypatch.delenv("CHEM_MGR_WEB_TOKEN", raising=False)
    app = create_app()
    app.dependency_overrides[get_manager] = lambda: manager
    return TestClient(app)


def _new_chem(**fields: Any) -> Dict[str, Any]:
    """
    功能:
        构造一份默认满足"必须有 substance"约束的化学品请求体.
    """
    payload: Dict[str, Any] = {
        "substance": "乙醇",
        "substance_english_name": "Ethanol",
        "physical_state": "liquid",
        "density": 0.789,
    }
    payload.update(fields)
    return payload


def test_health(client: TestClient) -> None:
    """功能: /api/health 返回 ok."""
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_create_get_update_delete(client: TestClient, manager: ChemicalManager) -> None:
    """功能: 化学品 CRUD 闭环."""
    # 创建
    r = client.post("/api/chemicals", json=_new_chem())
    assert r.status_code == 201, r.text
    body = r.json()
    row_id = body["id"]
    assert body["substance"] == "乙醇"
    assert body["density"] == pytest.approx(0.789)

    # 读取
    r = client.get(f"/api/chemicals/{row_id}")
    assert r.status_code == 200
    assert r.json()["substance_english_name"] == "Ethanol"

    # 更新
    r = client.put(f"/api/chemicals/{row_id}", json={"storage_location": "TB-2-5"})
    assert r.status_code == 200
    assert r.json()["storage_location"] == "TB-2-5"

    # 列表(带搜索)
    r = client.get("/api/chemicals", params={"q": "乙醇", "query_type": "name"})
    assert r.status_code == 200
    listing = r.json()
    assert listing["total"] == 1
    assert listing["items"][0]["substance"] == "乙醇"

    # 删除
    r = client.delete(f"/api/chemicals/{row_id}")
    assert r.status_code == 200
    assert r.json()["deleted"] is True

    # 删除后再 GET 应该 404
    r = client.get(f"/api/chemicals/{row_id}")
    assert r.status_code == 404


def test_create_substance_conflict_returns_409(client: TestClient) -> None:
    """功能: 重复 substance 写入触发 UNIQUE 约束 → 409."""
    r1 = client.post("/api/chemicals", json=_new_chem())
    assert r1.status_code == 201
    r2 = client.post("/api/chemicals", json=_new_chem())
    assert r2.status_code == 409, r2.text


def test_create_missing_substance_returns_400(client: TestClient) -> None:
    """功能: 缺少 substance 与英文名时返回 400."""
    r = client.post("/api/chemicals", json={"density": 1.0})
    assert r.status_code == 400


def test_get_unknown_returns_404(client: TestClient) -> None:
    """功能: 不存在的 id 返回 404."""
    r = client.get("/api/chemicals/999999")
    assert r.status_code == 404


def test_get_by_substance_returns_exact_match(client: TestClient) -> None:
    """功能: /api/chemicals/by-substance 按 substance 精确命中并返回 smiles."""
    create_resp = client.post(
        "/api/chemicals",
        json=_new_chem(substance="1,3,5-三异丙基苯", smiles="CC(C)C1=CC(C(C)C)=CC(C(C)C)=C1"),
    )
    assert create_resp.status_code == 201

    response = client.get(
        "/api/chemicals/by-substance",
        params={"substance": "1,3,5-三异丙基苯"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["id"] == create_resp.json()["id"]
    assert body["smiles"] == "CC(C)C1=CC(C(C)C)=CC(C(C)C)=C1"


def test_get_by_substance_returns_404_when_missing(client: TestClient) -> None:
    """功能: /api/chemicals/by-substance 未命中时返回 404."""
    response = client.get(
        "/api/chemicals/by-substance",
        params={"substance": "不存在的化学品"},
    )
    assert response.status_code == 404


def test_list_pagination(client: TestClient) -> None:
    """功能: 列表分页字段正确."""
    for name in ["甲醇", "乙醇", "丙醇"]:
        client.post("/api/chemicals", json=_new_chem(substance=name))

    r = client.get("/api/chemicals", params={"page": 1, "page_size": 2})
    body = r.json()
    assert body["total"] == 3
    assert body["page"] == 1
    assert body["page_size"] == 2
    assert len(body["items"]) == 2

    r2 = client.get("/api/chemicals", params={"page": 2, "page_size": 2})
    assert len(r2.json()["items"]) == 1


def test_integrity_endpoint(client: TestClient) -> None:
    """功能: /api/integrity 返回完整性报告."""
    client.post("/api/chemicals", json=_new_chem(cas_number="64-17-5"))
    r = client.get("/api/integrity")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    assert body["no_cas"] == 0


def test_export_csv(client: TestClient) -> None:
    """功能: /api/export.csv 导出 UTF-8-SIG CSV."""
    client.post("/api/chemicals", json=_new_chem())
    r = client.get("/api/export.csv")
    assert r.status_code == 200
    assert "csv" in r.headers["content-type"]
    text = r.content.decode("utf-8-sig")
    assert "乙醇" in text


def test_export_xlsx(client: TestClient) -> None:
    """功能: /api/export.xlsx 导出 xlsx, 响应 PK 魔数且可被 openpyxl 解析."""
    client.post("/api/chemicals", json=_new_chem(cas_number="64-17-5"))
    r = client.get("/api/export.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml.sheet" in r.headers["content-type"]
    # xlsx 本质为 zip, 以 PK\x03\x04 开头
    assert r.content[:4] == b"PK\x03\x04"

    # 进一步用 openpyxl 解析, 确认表头与数据行
    wb = openpyxl.load_workbook(filename=io.BytesIO(r.content), data_only=True)
    try:
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert "substance" in headers
        assert "cas_number" in headers
        # 第二行应该为刚插入的乙醇
        substance_idx = headers.index("substance")
        cas_idx = headers.index("cas_number")
        second_row = [cell.value for cell in ws[2]]
        assert second_row[substance_idx] == "乙醇"
        assert second_row[cas_idx] == "64-17-5"
    finally:
        wb.close()


def _build_xlsx_bytes(rows: list) -> bytes:
    """
    功能:
        用给定的行数据构建一个内存 xlsx 并返回字节流.
    参数:
        rows: List[Dict[str, Any]], 第一行的 keys 作为表头, 后续按该顺序填充.
    返回:
        bytes, xlsx 文件内容.
    """
    assert len(rows) > 0
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def _build_csv_bytes(rows: list) -> bytes:
    """
    功能:
        用给定的行数据构建 utf-8-sig 编码的 csv 文件字节流.
    参数:
        rows: List[Dict[str, Any]], 第一行的 keys 作为表头.
    返回:
        bytes, csv 文件内容.
    """
    assert len(rows) > 0
    headers = list(rows[0].keys())
    text_buf = io.StringIO(newline="")
    writer = csv.DictWriter(text_buf, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return text_buf.getvalue().encode("utf-8-sig")


def test_import_xlsx_round_trip(client: TestClient, manager: ChemicalManager) -> None:
    """功能: 上传 xlsx 后新化学品应出现在库中."""
    xlsx_bytes = _build_xlsx_bytes([
        {"substance": "甲醇", "substance_english_name": "Methanol", "cas_number": "67-56-1"},
        {"substance": "丙酮", "substance_english_name": "Acetone", "cas_number": "67-64-1"},
    ])
    r = client.post(
        "/api/import",
        files={"file": ("sample.xlsx", xlsx_bytes, "application/octet-stream")},
        data={"dry_run": "false"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["migrated"] == 2
    assert body["failed"] == 0
    assert manager.db.count() == 2

    # 再查一下 /api/chemicals 确认入库
    r2 = client.get("/api/chemicals", params={"q": "甲醇", "query_type": "name"})
    assert r2.status_code == 200
    assert r2.json()["total"] == 1


def test_import_csv_round_trip(client: TestClient, manager: ChemicalManager) -> None:
    """功能: 上传 csv 后新化学品应出现在库中."""
    csv_bytes = _build_csv_bytes([
        {"substance": "丙酮", "substance_english_name": "Acetone", "cas_number": "67-64-1"},
    ])
    r = client.post(
        "/api/import",
        files={"file": ("sample.csv", csv_bytes, "text/csv")},
        data={"dry_run": "false"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["migrated"] == 1
    assert body["failed"] == 0
    assert manager.db.count() == 1


def test_import_rejects_invalid_extension(client: TestClient) -> None:
    """功能: 上传非 xlsx/csv 应返回 400."""
    content = "substance,cas_number\nEthanol,64-17-5\n".encode("utf-8")
    r = client.post(
        "/api/import",
        files={"file": ("sample.txt", content, "text/plain")},
        data={"dry_run": "false"},
    )
    assert r.status_code == 400


def test_import_dry_run_does_not_write(client: TestClient, manager: ChemicalManager) -> None:
    """功能: dry_run=true 时不写入数据库."""
    baseline = manager.db.count()
    xlsx_bytes = _build_xlsx_bytes([
        {"substance": "四氢呋喃", "substance_english_name": "Tetrahydrofuran", "cas_number": "109-99-9"},
    ])
    r = client.post(
        "/api/import",
        files={"file": ("sample.xlsx", xlsx_bytes, "application/octet-stream")},
        data={"dry_run": "true"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["migrated"] == 1
    assert body["skipped"] == 0
    # 库内总数不变
    assert manager.db.count() == baseline


def test_import_skips_duplicates(client: TestClient, manager: ChemicalManager) -> None:
    """功能: 同一文件第二次上传, 全部计入 skipped."""
    xlsx_bytes = _build_xlsx_bytes([
        {"substance": "乙醚", "substance_english_name": "Diethyl ether", "cas_number": "60-29-7"},
    ])
    first = client.post(
        "/api/import",
        files={"file": ("dup.xlsx", xlsx_bytes, "application/octet-stream")},
        data={"dry_run": "false"},
    ).json()
    assert first["migrated"] == 1

    second = client.post(
        "/api/import",
        files={"file": ("dup.xlsx", xlsx_bytes, "application/octet-stream")},
        data={"dry_run": "false"},
    ).json()
    assert second["migrated"] == 0
    assert second["skipped"] == 1
    assert manager.db.count() == 1


def test_lookup_route_uses_manager(
    client: TestClient,
    manager: ChemicalManager,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """功能: /api/lookup 调用 manager.lookup_preview, 用 mock 避免发真实 HTTP."""
    captured: Dict[str, Any] = {}

    def fake_lookup(self, query: str, query_type: str):
        captured["query"] = query
        captured["query_type"] = query_type
        return {
            "row_data": {"substance": "乙醇", "cas_number": "64-17-5"},
            "chemicalbook_status": "ok",
            "chemicalbook_record_path": "",
        }

    monkeypatch.setattr(ChemicalManager, "lookup_preview", fake_lookup)
    r = client.post("/api/lookup", json={"query": "64-17-5", "query_type": "cas"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["success"] is True
    assert body["row_data"]["substance"] == "乙醇"
    assert body["chemicalbook_status"] == "ok"
    assert captured == {"query": "64-17-5", "query_type": "cas"}


def test_lookup_route_handles_empty_result(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """功能: 查询未命中时路由返回 success=False 并附带 message."""

    def fake_lookup(self, query: str, query_type: str):
        return None

    monkeypatch.setattr(ChemicalManager, "lookup_preview", fake_lookup)
    r = client.post("/api/lookup", json={"query": "unknown", "query_type": "name"})
    assert r.status_code == 200
    body = r.json()
    assert body["success"] is False
    assert body["message"]


def test_token_required_when_env_set(
    manager: ChemicalManager, monkeypatch: pytest.MonkeyPatch
) -> None:
    """功能: 设置 CHEM_MGR_WEB_TOKEN 后所有受保护路由必须带匹配 token."""
    monkeypatch.setenv("CHEM_MGR_WEB_TOKEN", "secret123")
    app = create_app()
    app.dependency_overrides[get_manager] = lambda: manager
    cli = TestClient(app)

    # 无 token → 401
    r = cli.get("/api/chemicals")
    assert r.status_code == 401

    # 错误 token → 401
    r = cli.get("/api/chemicals", headers={"X-API-Token": "wrong"})
    assert r.status_code == 401

    # 正确 token → 200
    r = cli.get("/api/chemicals", headers={"X-API-Token": "secret123"})
    assert r.status_code == 200

    # 健康探针不受鉴权影响
    r = cli.get("/api/health")
    assert r.status_code == 200
