# -*- coding: utf-8 -*-
"""
功能:
    从 xlsx 或 csv 文件向化学品库追加化学品.
    统一完成表头读取, 列名适配, 数值类型转换, 重复检查和插入.
    不做源文件重命名或备份, 仅完成入库并返回统计.
"""

import csv
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from ..config.constants import NUMERIC_COLUMNS
from ..config.setting import Settings
from ..controller.chemical_db import ChemicalDB

logger = logging.getLogger("ChemicalImporter")

__all__ = ["import_to_library"]


def _build_header_map(worksheet: Any) -> Dict[str, int]:
    """
    功能:
        从 Excel 工作表首行构建表头名到列号映射.
    参数:
        worksheet: openpyxl 工作表对象.
    返回:
        Dict[str, int], 表头名 -> 列号.
    """
    header_map: Dict[str, int] = {}
    for col_idx in range(1, worksheet.max_column + 1):
        header_val = worksheet.cell(row=1, column=col_idx).value
        if header_val is None:
            continue
        header_map[str(header_val).strip()] = col_idx
    return header_map


def _extract_row_xlsx(
    worksheet: Any,
    header_map: Dict[str, int],
    row_index: int,
) -> Dict[str, Any]:
    """
    功能:
        从工作表中提取单行数据为字典, 空单元格忽略.
    参数:
        worksheet: openpyxl 工作表对象.
        header_map: Dict[str, int], 表头映射.
        row_index: int, 行号 (1-based).
    返回:
        Dict[str, Any], 字段字典.
    """
    row_data: Dict[str, Any] = {}
    for column_name, column_index in header_map.items():
        val = worksheet.cell(row=row_index, column=column_index).value
        if val is not None:
            row_data[column_name] = val
    return row_data


def _normalize_row(row_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    功能:
        标准化行数据: 统一中文名字段, 适配密度等列名, 数值列转 float, 字符串去空白.
    参数:
        row_data: Dict[str, Any], 原始行数据.
    返回:
        Dict[str, Any], 标准化后的行数据.
    """
    result = dict(row_data)

    # substance / substance_chinese_name 统一
    substance = str(result.get("substance") or "").strip()
    chinese_name = str(result.get("substance_chinese_name") or "").strip()
    if substance == "" and chinese_name != "":
        result["substance"] = chinese_name
    if "substance_chinese_name" in result:
        del result["substance_chinese_name"]

    # density (g/mL) 列名适配
    if "density (g/mL)" in result:
        result["density"] = result.pop("density (g/mL)")

    # active_content(mol/L or wt%) 列名适配
    if "active_content(mol/L or wt%)" in result:
        result["active_content"] = result.pop("active_content(mol/L or wt%)")

    # 数值列转换
    for col in NUMERIC_COLUMNS:
        if col in result and result[col] is not None:
            try:
                result[col] = float(result[col])
            except (TypeError, ValueError):
                result[col] = None

    # 字符串列去空白
    for key, val in list(result.items()):
        if isinstance(val, str):
            result[key] = val.strip()

    return result


def _read_rows_from_xlsx(file_path: Path) -> List[Dict[str, Any]]:
    """
    功能:
        读取 xlsx 文件并返回原始行数据列表.
    参数:
        file_path: Path, xlsx 文件路径.
    返回:
        List[Dict[str, Any]], 每行一个字典, 空行被忽略.
    异常:
        ValueError: 首行无表头时抛出.
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    try:
        ws = wb.active
        header_map = _build_header_map(ws)
        if len(header_map) == 0:
            raise ValueError("Excel 首行无表头, 无法导入")

        rows: List[Dict[str, Any]] = []
        for row_idx in range(2, ws.max_row + 1):
            raw = _extract_row_xlsx(ws, header_map, row_idx)
            if len(raw) == 0:
                continue
            rows.append(raw)
        return rows
    finally:
        wb.close()


def _read_rows_from_csv(file_path: Path) -> List[Dict[str, Any]]:
    """
    功能:
        读取 csv 文件并返回原始行数据列表.
        采用 utf-8-sig 编码兼容 Excel 导出的带 BOM 文件.
        跳过所有字段为空的行.
    参数:
        file_path: Path, csv 文件路径.
    返回:
        List[Dict[str, Any]], 每行一个字典, 空行被忽略.
    异常:
        ValueError: 首行无表头时抛出.
    """
    rows: List[Dict[str, Any]] = []
    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None or len([h for h in reader.fieldnames if h]) == 0:
            raise ValueError("CSV 首行无表头, 无法导入")

        for raw in reader:
            # 剔除空值并去掉 DictReader 对缺列补 None 的键, 保证下游判断一致
            cleaned: Dict[str, Any] = {}
            for key, val in raw.items():
                if key is None:
                    continue
                if val is None:
                    continue
                if isinstance(val, str) and val.strip() == "":
                    continue
                cleaned[str(key).strip()] = val
            if len(cleaned) == 0:
                continue
            rows.append(cleaned)
    return rows


def _read_rows_by_extension(file_path: Path) -> List[Dict[str, Any]]:
    """
    功能:
        根据文件后缀分派到 xlsx 或 csv 读取器.
    参数:
        file_path: Path, 文件路径.
    返回:
        List[Dict[str, Any]], 原始行数据.
    异常:
        ValueError: 后缀不受支持时抛出.
        FileNotFoundError: 文件不存在时抛出.
    """
    if file_path.exists() is False:
        raise FileNotFoundError(f"导入文件不存在: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix == ".xlsx":
        return _read_rows_from_xlsx(file_path)
    if suffix == ".csv":
        return _read_rows_from_csv(file_path)
    raise ValueError(f"不支持的文件格式: {suffix}, 仅支持 .xlsx 或 .csv")


def import_to_library(
    file_path: str,
    db_path: Optional[str] = None,
    dry_run: bool = False,
) -> Dict[str, int]:
    """
    功能:
        从 xlsx 或 csv 文件向化学品库追加化学品.
        统一经过列名适配, 重复检查, 插入流程.
    参数:
        file_path: str, 源文件路径, 以 .xlsx 或 .csv 结尾.
        db_path: Optional[str], SQLite 文件路径, None 时使用默认配置.
        dry_run: bool, True 时仅统计有效行数不写入.
    返回:
        Dict[str, int], 包含 migrated(成功入库), skipped(重复或无标识字段), failed(写入失败) 三项计数.
    """
    source = Path(file_path)

    settings = Settings.from_env()
    target_db_path = Path(db_path) if db_path is not None else settings.db_path

    logger.info("开始从文件导入化学品: %s -> %s (dry_run=%s)", source, target_db_path, dry_run)

    raw_rows = _read_rows_by_extension(source)

    stats = {"migrated": 0, "skipped": 0, "failed": 0}

    if dry_run is True:
        # 预览模式下仅统计包含核心标识字段的有效行
        for raw in raw_rows:
            normalized = _normalize_row(raw)
            if _has_core_identifier(normalized) is False:
                stats["skipped"] += 1
                continue
            stats["migrated"] += 1
        logger.info("预览完成, 预计导入 %d 行, 跳过 %d 行", stats["migrated"], stats["skipped"])
        return stats

    db = ChemicalDB(target_db_path)
    try:
        for row_idx, raw in enumerate(raw_rows, start=2):
            normalized = _normalize_row(raw)

            if _has_core_identifier(normalized) is False:
                logger.warning("跳过无标识字段的行: row=%d", row_idx)
                stats["skipped"] += 1
                continue

            # 重复检查
            dup = db.find_duplicate(normalized)
            if dup is not None:
                label, val, existing_id = dup
                logger.debug("跳过重复行: row=%d, %s=%s, existing_id=%d", row_idx, label, val, existing_id)
                stats["skipped"] += 1
                continue

            try:
                db.insert_row(normalized)
                stats["migrated"] += 1
            except Exception as exc:
                logger.error("插入失败: row=%d, err=%s", row_idx, exc)
                stats["failed"] += 1
    finally:
        db.close()

    logger.info(
        "导入完成: migrated=%d, skipped=%d, failed=%d",
        stats["migrated"], stats["skipped"], stats["failed"],
    )

    return stats


def _has_core_identifier(row_data: Dict[str, Any]) -> bool:
    """
    功能:
        判断行数据是否至少包含一个可识别化合物的核心字段.
    参数:
        row_data: Dict[str, Any], 标准化后的行数据.
    返回:
        bool, True 表示至少包含 CAS, 英文名, 中文名中的一个.
    """
    return any([
        str(row_data.get("cas_number") or "").strip() != "",
        str(row_data.get("substance_english_name") or "").strip() != "",
        str(row_data.get("substance") or "").strip() != "",
    ])
