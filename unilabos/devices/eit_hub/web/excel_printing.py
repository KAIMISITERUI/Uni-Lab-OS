# -*- coding: utf-8 -*-
"""
功能:
    提供 EIT Hub 上料表格打印能力.
    直接读取 batch_in_tray 工作表并通过 Windows GDI 绘制到指定本机打印机.
"""

from __future__ import annotations

import logging
import math
from pathlib import Path
from typing import Any, List, Sequence, Tuple

from openpyxl import load_workbook

logger = logging.getLogger("EITHubExcelPrinting")

DEFAULT_BATCH_IN_TABLE_PRINTER = "HP Laser MFP 1136-1139 1188"

_BATCH_IN_SHEET_NAME = "batch_in_tray"
_BATCH_IN_COLUMNS = 5
_COL_WIDTH_RATIOS = (0.11, 0.24, 0.31, 0.12, 0.22)
_PRINTER_ATTRIBUTE_WORK_OFFLINE = 0x00000400
_PRINTER_STATUS_OFFLINE = 0x00000080
_DM_ORIENTATION = 0x00000001
_DMORIENT_LANDSCAPE = 2
_GB2312_CHARSET = 134
_ANTIALIASED_QUALITY = 4


def print_batch_in_table(
    workbook_path: Path,
    printer_name: str = DEFAULT_BATCH_IN_TABLE_PRINTER,
) -> dict[str, Any]:
    """
    功能:
        读取 batch_in_tray 工作表并直接打印到指定 Windows 打印机.
    参数:
        workbook_path: Path, 上料 Excel 文件路径.
        printer_name: str, Windows 打印机名称.
    返回:
        dict[str, Any], 包含打印机名称, 文件路径和打印行数.
    """
    target_path = Path(workbook_path).resolve()
    if target_path.exists() is False:
        raise FileNotFoundError(f"未找到上料文件: {target_path}")

    rows = _read_batch_in_print_rows(target_path)
    win32print, win32ui, win32con, win32gui = _load_windows_print_modules()
    _ensure_printer_ready(win32print, printer_name)

    printer_dc = _create_printer_dc(win32print, win32ui, win32con, win32gui, printer_name)
    try:
        _draw_rows_to_printer(printer_dc, win32con, win32gui, rows, target_path.name)
    finally:
        printer_dc.DeleteDC()

    printed_rows = max(len(rows) - 1, 0)
    logger.info("上料表格打印任务已提交 | 文件: %s | 打印机: %s | 数据行数: %d", target_path, printer_name, printed_rows)
    return {
        "printer_name": printer_name,
        "file_path": str(target_path),
        "printed_rows": printed_rows,
    }


def _load_windows_print_modules() -> Tuple[Any, Any, Any, Any]:
    """
    功能:
        延迟导入 Windows 打印模块.
    参数:
        无.
    返回:
        Tuple[Any, Any, Any, Any], win32print, win32ui, win32con, win32gui 模块.
    """
    try:
        import win32con
        import win32gui
        import win32print
        import win32ui
    except ImportError as exc:
        raise RuntimeError("当前 Python 环境缺少 pywin32, 无法打印上料表格.") from exc

    return win32print, win32ui, win32con, win32gui


def _read_batch_in_print_rows(workbook_path: Path) -> List[List[str]]:
    """
    功能:
        从上料 Excel 中读取 A:E 打印区域.
    参数:
        workbook_path: Path, 上料 Excel 文件路径.
    返回:
        List[List[str]], 表头和数据行.
    """
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        worksheet = _select_batch_in_sheet(workbook)
        last_row = _find_last_non_empty_row(worksheet)
        rows: List[List[str]] = []
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=last_row,
            min_col=1,
            max_col=_BATCH_IN_COLUMNS,
            values_only=True,
        ):
            rows.append([_cell_text(value) for value in row])
        if len(rows) == 0:
            return [["position", "tray_type", "content", "shelf_position", "storage"]]
        return rows
    finally:
        workbook.close()


def _select_batch_in_sheet(workbook: Any) -> Any:
    """
    功能:
        定位 batch_in_tray 工作表.
    参数:
        workbook: Any, openpyxl 工作簿对象.
    返回:
        Any, openpyxl 工作表对象.
    """
    for sheet_name in workbook.sheetnames:
        if str(sheet_name).strip().lower() == _BATCH_IN_SHEET_NAME:
            return workbook[sheet_name]
    raise RuntimeError("上料文件缺少 batch_in_tray 工作表, 已取消打印.")


def _find_last_non_empty_row(worksheet: Any) -> int:
    """
    功能:
        查找 A:E 范围内最后一个非空行.
    参数:
        worksheet: Any, openpyxl 工作表对象.
    返回:
        int, 最后一个非空行号, 最小为 1.
    """
    last_row = 1
    for row_index, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=_BATCH_IN_COLUMNS,
            values_only=True,
        ),
        start=1,
    ):
        if any(_cell_text(value) != "" for value in row) is True:
            last_row = row_index
    return last_row


def _cell_text(value: Any) -> str:
    """
    功能:
        将单元格值转换为打印文本.
    参数:
        value: Any, 单元格原始值.
    返回:
        str, 去空格后的文本.
    """
    if value is None:
        return ""
    return str(value).strip()


def _ensure_printer_ready(win32print: Any, printer_name: str) -> None:
    """
    功能:
        校验指定 Windows 打印机存在且未离线.
    参数:
        win32print: Any, pywin32 打印模块.
        printer_name: str, Windows 打印机名称.
    返回:
        None.
    """
    printer_info = None
    flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
    for item in win32print.EnumPrinters(flags, None, 2):
        name = _read_printer_field(item, "pPrinterName", 1)
        if name == printer_name:
            printer_info = item
            break

    if printer_info is None:
        raise RuntimeError(f"未找到上料表格打印机: {printer_name}")

    attributes = int(_read_printer_field(printer_info, "Attributes", 13) or 0)
    status = int(_read_printer_field(printer_info, "Status", 18) or 0)
    if (attributes & _PRINTER_ATTRIBUTE_WORK_OFFLINE) != 0:
        raise RuntimeError(f"上料表格打印机处于离线状态: {printer_name}")
    if (status & _PRINTER_STATUS_OFFLINE) != 0:
        raise RuntimeError(f"上料表格打印机处于离线状态: {printer_name}")


def _read_printer_field(printer_info: Any, key: str, tuple_index: int) -> Any:
    """
    功能:
        兼容 dict 和 tuple 两种 pywin32 打印机信息结构.
    参数:
        printer_info: Any, EnumPrinters 返回项.
        key: str, dict 字段名.
        tuple_index: int, tuple 字段下标.
    返回:
        Any, 字段值.
    """
    if isinstance(printer_info, dict):
        return printer_info.get(key)
    if isinstance(printer_info, tuple) and len(printer_info) > tuple_index:
        return printer_info[tuple_index]
    return None


def _create_printer_dc(
    win32print: Any,
    win32ui: Any,
    win32con: Any,
    win32gui: Any,
    printer_name: str,
) -> Any:
    """
    功能:
        创建指向指定打印机的 GDI 设备上下文, 优先设置横向打印.
    参数:
        win32print: Any, pywin32 打印模块.
        win32ui: Any, pywin32 UI 模块.
        win32con: Any, pywin32 常量模块.
        win32gui: Any, pywin32 GUI 模块.
        printer_name: str, Windows 打印机名称.
    返回:
        Any, win32ui 打印设备上下文.
    """
    printer_handle = win32print.OpenPrinter(printer_name)
    try:
        printer_info = win32print.GetPrinter(printer_handle, 2)
        dev_mode = printer_info.get("pDevMode") if isinstance(printer_info, dict) else None
        if dev_mode is not None:
            dev_mode.Orientation = getattr(win32con, "DMORIENT_LANDSCAPE", _DMORIENT_LANDSCAPE)
            dev_mode.Fields = int(dev_mode.Fields) | getattr(win32con, "DM_ORIENTATION", _DM_ORIENTATION)
        raw_dc = win32gui.CreateDC("WINSPOOL", printer_name, dev_mode)
        return win32ui.CreateDCFromHandle(raw_dc)
    finally:
        win32print.ClosePrinter(printer_handle)


def _draw_rows_to_printer(
    dc: Any,
    win32con: Any,
    win32gui: Any,
    rows: Sequence[Sequence[str]],
    document_name: str,
) -> None:
    """
    功能:
        将上料表格绘制到打印机设备上下文.
    参数:
        dc: Any, win32ui 打印设备上下文.
        win32con: Any, pywin32 常量模块.
        win32gui: Any, pywin32 GUI 模块.
        rows: Sequence[Sequence[str]], 表头和数据行.
        document_name: str, 打印任务名称.
    返回:
        None.
    """
    dpi_x = dc.GetDeviceCaps(getattr(win32con, "LOGPIXELSX", 88))
    dpi_y = dc.GetDeviceCaps(getattr(win32con, "LOGPIXELSY", 90))
    page_width = dc.GetDeviceCaps(getattr(win32con, "HORZRES", 8))
    page_height = dc.GetDeviceCaps(getattr(win32con, "VERTRES", 10))
    margin_x = _mm_to_device_units(10.0, dpi_x)
    margin_y = _mm_to_device_units(10.0, dpi_y)
    usable_width = max(page_width - margin_x * 2, 1)
    table_right = margin_x + usable_width
    page_bottom = page_height - margin_y
    col_widths = _build_column_widths(usable_width)

    title_font = _create_font(dc, "Microsoft YaHei", 14, dpi_y, bold=True)
    header_font = _create_font(dc, "Microsoft YaHei", 9, dpi_y, bold=True)
    text_font = _create_font(dc, "Microsoft YaHei", 9, dpi_y, bold=False)
    line_height = _point_to_device_units(12, dpi_y)
    padding = max(_mm_to_device_units(1.5, dpi_x), 4)
    title_height = _point_to_device_units(22, dpi_y)
    header_height = max(_point_to_device_units(24, dpi_y), line_height + padding * 2)

    header_row = list(rows[0]) if len(rows) > 0 else ["position", "tray_type", "content", "shelf_position", "storage"]
    data_rows = list(rows[1:]) if len(rows) > 1 else []
    data_index = 0
    doc_started = False

    try:
        dc.StartDoc(f"上料表格 - {document_name}")
        doc_started = True
        while True:
            dc.StartPage()
            dc.SetBkMode(getattr(win32con, "TRANSPARENT", 1))
            y = margin_y

            dc.SelectObject(title_font)
            title_rect = (margin_x, y, table_right, y + title_height)
            _draw_text(
                dc,
                win32gui,
                win32con,
                "上料表格",
                title_rect,
                getattr(win32con, "DT_LEFT", 0) | getattr(win32con, "DT_VCENTER", 4),
            )
            y += title_height

            dc.SelectObject(header_font)
            _draw_table_row(dc, win32gui, win32con, header_row, margin_x, y, col_widths, header_height, padding)
            y += header_height

            dc.SelectObject(text_font)
            if len(data_rows) == 0:
                dc.EndPage()
                break

            while data_index < len(data_rows):
                row = data_rows[data_index]
                row_height = _estimate_row_height(row, col_widths, line_height, padding)
                if y + row_height > page_bottom and y > margin_y + title_height + header_height:
                    break
                _draw_table_row(dc, win32gui, win32con, row, margin_x, y, col_widths, row_height, padding)
                y += row_height
                data_index += 1

            dc.EndPage()
            if data_index >= len(data_rows):
                break
        dc.EndDoc()
    except Exception:
        if doc_started is True:
            try:
                dc.AbortDoc()
            except Exception:
                logger.debug("取消打印任务失败, 已忽略.")
        raise


def _create_font(dc: Any, font_name: str, point_size: int, dpi_y: int, bold: bool = False) -> Any:
    """
    功能:
        创建打印字体.
    参数:
        dc: Any, win32ui 打印设备上下文.
        font_name: str, 字体名称.
        point_size: int, 字号.
        dpi_y: int, 打印机纵向 DPI.
        bold: bool, 是否加粗.
    返回:
        Any, win32ui 字体对象.
    """
    win32ui = __import__("win32ui")
    return win32ui.CreateFont(
        {
            "name": font_name,
            "height": -_point_to_device_units(point_size, dpi_y),
            "weight": 700 if bold is True else 400,
            "charset": _GB2312_CHARSET,
            "quality": _ANTIALIASED_QUALITY,
        }
    )


def _build_column_widths(usable_width: int) -> List[int]:
    """
    功能:
        根据页面宽度计算上料表格各列宽度.
    参数:
        usable_width: int, 可用打印宽度.
    返回:
        List[int], 每列宽度.
    """
    widths = [int(usable_width * ratio) for ratio in _COL_WIDTH_RATIOS]
    diff = usable_width - sum(widths)
    widths[-1] += diff
    return widths


def _draw_table_row(
    dc: Any,
    win32gui: Any,
    win32con: Any,
    row: Sequence[str],
    left: int,
    top: int,
    col_widths: Sequence[int],
    row_height: int,
    padding: int,
) -> None:
    """
    功能:
        绘制一行表格.
    参数:
        dc: Any, win32ui 打印设备上下文.
        win32gui: Any, pywin32 GUI 模块.
        win32con: Any, pywin32 常量模块.
        row: Sequence[str], 行数据.
        left: int, 左边界.
        top: int, 上边界.
        col_widths: Sequence[int], 列宽.
        row_height: int, 行高.
        padding: int, 单元格内边距.
    返回:
        None.
    """
    text_flags = (
        getattr(win32con, "DT_LEFT", 0)
        | getattr(win32con, "DT_TOP", 0)
        | getattr(win32con, "DT_WORDBREAK", 0x10)
        | getattr(win32con, "DT_NOPREFIX", 0x800)
    )
    x = left
    for col_index, col_width in enumerate(col_widths):
        cell_text = row[col_index] if col_index < len(row) else ""
        rect = (x, top, x + col_width, top + row_height)
        dc.Rectangle(rect)
        text_rect = (rect[0] + padding, rect[1] + padding, rect[2] - padding, rect[3] - padding)
        _draw_text(dc, win32gui, win32con, str(cell_text), text_rect, text_flags)
        x += col_width


def _draw_text(dc: Any, win32gui: Any, win32con: Any, text: str, rect: Tuple[int, int, int, int], flags: int) -> None:
    """
    功能:
        使用 Unicode GDI 绘制文本, 避免中文经过 ANSI 路径后乱码.
    参数:
        dc: Any, win32ui 打印设备上下文.
        win32gui: Any, pywin32 GUI 模块.
        win32con: Any, pywin32 常量模块.
        text: str, 待绘制文本.
        rect: Tuple[int, int, int, int], 绘制区域.
        flags: int, DrawText 标志.
    返回:
        None.
    """
    draw_flags = flags | getattr(win32con, "DT_NOPREFIX", 0x800)
    content = str(text)
    win32gui.DrawText(dc.GetSafeHdc(), content, len(content), rect, draw_flags)


def _estimate_row_height(row: Sequence[str], col_widths: Sequence[int], line_height: int, padding: int) -> int:
    """
    功能:
        根据文本长度估算打印行高.
    参数:
        row: Sequence[str], 行数据.
        col_widths: Sequence[int], 列宽.
        line_height: int, 单行高度.
        padding: int, 单元格内边距.
    返回:
        int, 行高.
    """
    max_lines = 1
    for col_index, col_width in enumerate(col_widths):
        text = row[col_index] if col_index < len(row) else ""
        chars_per_line = max(int((col_width - padding * 2) / max(line_height * 0.55, 1)), 6)
        line_count = max(math.ceil(len(str(text)) / chars_per_line), 1)
        max_lines = max(max_lines, min(line_count, 5))
    return max(line_height * max_lines + padding * 2, line_height + padding * 2)


def _point_to_device_units(point_size: float, dpi: int) -> int:
    """
    功能:
        将磅值转换为打印机设备单位.
    参数:
        point_size: float, 字号或长度, 单位 pt.
        dpi: int, 打印机 DPI.
    返回:
        int, 设备单位.
    """
    return max(int(point_size * dpi / 72), 1)


def _mm_to_device_units(mm_value: float, dpi: int) -> int:
    """
    功能:
        将毫米转换为打印机设备单位.
    参数:
        mm_value: float, 毫米值.
        dpi: int, 打印机 DPI.
    返回:
        int, 设备单位.
    """
    return max(int(mm_value * dpi / 25.4), 1)
