# -*- coding: utf-8 -*-
"""
功能:
    覆盖 EIT Hub 合成工站 Excel 编解码行为.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List
from unittest.mock import patch

import openpyxl

from unilabos.devices.eit_analysis_station.processor.yield_calculator import YieldCalculator
from unilabos.devices.eit_hub.web.excel_codec import (
    read_batch_in_template,
    read_reaction_template,
    write_batch_in_template,
    write_reaction_template,
)


def _create_template(path: Path) -> None:
    """
    功能:
        创建测试用合成任务模板.
    参数:
        path: Path, 输出路径.
    返回:
        None.
    """
    workbook = openpyxl.Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = "实验方案设定"
        worksheet["A1"] = "实验设定"
        worksheet["A2"] = "实验名称"
        worksheet["B2"] = "旧任务"
        worksheet["A3"] = "实验ID"
        worksheet["B3"] = 1
        worksheet["A4"] = "反应设定"
        worksheet["A5"] = "反应规模(mmol)"
        worksheet["B5"] = 0.2
        worksheet["A6"] = "反应器类型"
        worksheet["B6"] = "heat"
        worksheet["A7"] = "自动加磁子"
        worksheet["B7"] = "是"
        worksheet["A8"] = "内标设定"
        worksheet["A9"] = "内标种类"
        worksheet["B9"] = "1,3,5-三异丙基苯"
        worksheet["A10"] = "内标用量(μL/mg)"
        worksheet["B10"] = 10
        worksheet["A11"] = "分析方法设定"
        worksheet["A12"] = "GC_MS"
        worksheet["B12"] = "280_12min"
        worksheet["A13"] = "UPLC_QTOF"
        worksheet["B13"] = "Generic_5min"
        worksheet["A14"] = "HPLC"
        worksheet["C1"] = "实验编号"
        worksheet["D1"] = "试剂"
        worksheet["E1"] = "试剂量"
        worksheet["F1"] = "试剂"
        worksheet["G1"] = "试剂量"
        for index in range(1, 13):
            row_index = index + 1
            worksheet.cell(row_index, 3, index)
            worksheet.cell(row_index, 4, "乙腈")
            worksheet.cell(row_index, 5, "1mL")

        gc_sheet = workbook.create_sheet("GC产率计算")
        gc_sheet["A1"] = "内标SMILES"
        gc_sheet["B1"] = "CC"
        gc_sheet["A2"] = "内标预期RT(min)"
        gc_sheet["B2"] = "6.84"
        gc_sheet["A3"] = "产率计算方法"
        gc_sheet["B3"] = "ECN"
        gc_sheet["A4"] = "标准曲线斜率"
        gc_sheet["A5"] = "标准曲线截距"
        gc_sheet["A6"] = "响应因子"
        gc_sheet["A8"] = "适用实验"
        gc_sheet["B8"] = "目标产物名称"
        gc_sheet["C8"] = "当量(eq)"
        gc_sheet["D8"] = "SMILES"
        gc_sheet["E8"] = "预期RT(min)"
        gc_sheet["A9"] = "1-12"
        gc_sheet["B9"] = "对叔丁基苯甲醛"
        gc_sheet["C9"] = 1
        gc_sheet["D9"] = "O=CC1=CC=C(C(C)(C)C)C=C1"
        workbook.save(path)
    finally:
        workbook.close()


def _create_legacy_template_without_gc_sheet(path: Path) -> None:
    """
    功能:
        创建不包含 GC产率计算 Sheet 的旧版实验模板.
    参数:
        path: Path, 输出路径.
    返回:
        None.
    """
    workbook = openpyxl.Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = "实验方案设定"
        worksheet["A1"] = "实验设定"
        worksheet["A2"] = "实验名称"
        worksheet["B2"] = "旧版任务"
        worksheet["A3"] = "实验ID"
        worksheet["B3"] = 17
        worksheet["C1"] = "实验编号"
        worksheet["D1"] = "试剂"
        worksheet["E1"] = "试剂量"
        for index in range(1, 13):
            row_index = index + 1
            worksheet.cell(row=row_index, column=3, value=index)
            worksheet.cell(row=row_index, column=4, value="乙腈")
            worksheet.cell(row=row_index, column=5, value="1mL")
        workbook.save(path)
    finally:
        workbook.close()


def _updated_gc_ms_yield() -> Dict[str, Any]:
    """
    功能:
        返回测试用 GC-MS 产率配置 payload.
    返回:
        Dict[str, Any], GC-MS 产率配置.
    """
    return {
        "internal_standard_smiles": "CC(C)C1=CC(C(C)C)=CC(C(C)C)=C1",
        "internal_standard_expected_rt": 6.84,
        "yield_method": "标准曲线",
        "curve_slope": 1.25,
        "curve_intercept": 0.05,
        "response_factor": "",
        "products": [
            {
                "applicable_experiments": "1-6",
                "product_name": "对叔丁基苯甲醛",
                "equivalent": 1,
                "smiles": "O=CC1=CC=C(C(C)(C)C)C=C1",
                "expected_rt": 8.1,
            },
            {
                "applicable_experiments": "7-12",
                "product_name": "4-甲基苯甲醛",
                "equivalent": 1,
                "smiles": "CC1=CC=C(C=O)C=C1",
                "expected_rt": "",
            },
            {
                "applicable_experiments": "",
                "product_name": "",
                "equivalent": "",
                "smiles": "",
                "expected_rt": "",
            },
        ],
    }


def _updated_payload() -> Dict[str, Any]:
    """
    功能:
        返回测试用 Web 模板数据.
    返回:
        Dict[str, Any], 模板 payload.
    """
    rows: List[List[Any]] = []
    for index in range(1, 13):
        rows.append([index, "对叔丁基苯甲醛", f"{index}.0eq", "乙腈", "1mL"])
    return {
        "param_rows": [
            {"name": "实验设定", "value": "", "type": "section"},
            {"name": "实验名称", "value": "Web任务", "type": "parameter"},
            {"name": "实验ID", "value": 1, "type": "parameter"},
            {"name": "反应规模(mmol)", "value": 0.5, "type": "parameter"},
            {"name": "反应器类型", "value": "heat", "type": "parameter"},
            {"name": "自动加磁子", "value": "否", "type": "parameter"},
            {"name": "内标设定", "value": "", "type": "section"},
            {"name": "内标种类", "value": "1,3,5-三异丙基苯", "type": "parameter"},
            {"name": "内标用量(μL/mg)", "value": 10, "type": "parameter"},
            {"name": "分析方法设定", "value": "", "type": "section"},
            {"name": "GC_MS", "value": "280_12min(1-3)", "type": "parameter"},
            {"name": "UPLC_QTOF", "value": "Generic_5min", "type": "parameter"},
            {"name": "HPLC", "value": "", "type": "parameter"},
        ],
        "gc_ms_yield": _updated_gc_ms_yield(),
        "headers": ["实验编号", "试剂", "试剂量", "试剂", "试剂量"],
        "rows": rows,
    }


def _create_batch_in_template(path: Path) -> None:
    """
    功能:
        创建测试用上料文件.
    参数:
        path: Path, 输出路径.
    返回:
        None.
    """
    workbook = openpyxl.Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = "batch_in_tray"
        worksheet.append(["position", "tray_type", "content", "shelf_position", "storage"])
        worksheet.append(["TB-2-1", "2 mL试剂瓶托盘(201000705)", "A1|乙腈|1mL", "3-1", "乙腈|A柜"])
        options_sheet = workbook.create_sheet("validation_meta")
        options_sheet.cell(row=1, column=1, value="2 mL试剂瓶托盘(201000705) [A1-F8]")
        options_sheet.cell(row=2, column=1, value="30 mL粉桶托盘(201000710) [A1-B1]")
        workbook.save(path)
    finally:
        workbook.close()


def _updated_batch_in_payload() -> Dict[str, Any]:
    """
    功能:
        返回测试用上料表格 payload.
    返回:
        Dict[str, Any], 上料表格 payload.
    """
    return {
        "headers": ["position", "tray_type", "content", "shelf_position", "storage"],
        "rows": [
            ["TB-2-2", "30 mL粉桶托盘(201000710)", "A1|碳酸钾|100mg", "3-2", "碳酸钾|B柜"],
            ["TB-2-3", "2 mL反应试管托盘(201000726)", "12", "3-3", "2 mL反应管|耗材库"],
        ],
    }


def test_reaction_template_read_write_round_trip(tmp_path: Path) -> None:
    """
    功能:
        验证模板读取, Web 数据写回和再次读取保持一致.
    """
    template_path = tmp_path / "reaction_template.xlsx"
    _create_template(template_path)

    original = read_reaction_template(template_path)
    assert original["params"]["实验名称"] == "旧任务"
    assert original["params"]["GC_MS"] == "280_12min"
    assert original["params"]["UPLC_QTOF"] == "Generic_5min"
    assert original["params"]["HPLC"] == ""
    param_names = [item["name"] for item in original["param_rows"]]
    assert "分析方法设定" in param_names
    assert "GC_MS" in param_names
    assert "UPLC_QTOF" in param_names
    assert "HPLC" in param_names
    assert original["headers"] == ["实验编号", "试剂", "试剂量", "试剂", "试剂量"]
    assert len(original["rows"]) == 12
    assert original["has_gc_ms_yield_sheet"] is True
    assert original["gc_ms_yield"]["internal_standard_smiles"] == "CC"
    assert original["gc_ms_yield"]["yield_method"] == "ECN"
    assert original["gc_ms_yield"]["products"][0]["product_name"] == "对叔丁基苯甲醛"

    saved = write_reaction_template(_updated_payload(), template_path)

    assert saved["params"]["实验名称"] == "Web任务"
    assert saved["params"]["反应规模(mmol)"] == 0.5
    assert saved["params"]["GC_MS"] == "280_12min(1-3)"
    assert saved["params"]["UPLC_QTOF"] == "Generic_5min"
    assert saved["params"]["HPLC"] == ""
    assert saved["gc_ms_yield"]["internal_standard_smiles"] == _updated_gc_ms_yield()["internal_standard_smiles"]
    assert saved["gc_ms_yield"]["yield_method"] == "标准曲线"
    assert len(saved["gc_ms_yield"]["products"]) == 2
    assert saved["gc_ms_yield"]["products"][1]["product_name"] == "4-甲基苯甲醛"
    assert saved["rows"][0] == [1, "对叔丁基苯甲醛", "1.0eq", "乙腈", "1mL"]

    workbook = openpyxl.load_workbook(template_path, data_only=False)
    try:
        assert "GC产率计算" in workbook.sheetnames
        assert workbook["GC产率计算"]["B1"].value == _updated_gc_ms_yield()["internal_standard_smiles"]
        assert workbook["GC产率计算"]["B3"].value == "标准曲线"
        assert workbook["GC产率计算"]["B4"].value == 1.25
        assert workbook["GC产率计算"]["B5"].value == 0.05
        assert workbook["GC产率计算"]["A10"].value == "7-12"
        assert workbook["GC产率计算"]["B10"].value == "4-甲基苯甲醛"
        assert workbook["实验方案设定"]["B2"].value == "Web任务"
        assert workbook["实验方案设定"]["B9"].value == "1,3,5-三异丙基苯"
        assert workbook["实验方案设定"]["B10"].value == 10
        assert workbook["实验方案设定"]["B12"].value == "280_12min(1-3)"
        assert workbook["实验方案设定"]["B13"].value == "Generic_5min"
        assert workbook["实验方案设定"]["B14"].value is None
        assert workbook["实验方案设定"]["D2"].value == "对叔丁基苯甲醛"
    finally:
        workbook.close()


def test_reaction_template_write_result_can_be_parsed_by_yield_calculator(tmp_path: Path) -> None:
    """
    功能:
        验证 Web 保存后的 reaction_template 可被 YieldCalculator 正确解析 GC-MS 产率配置.
    """
    template_path = tmp_path / "reaction_template.xlsx"
    _create_template(template_path)
    write_reaction_template(_updated_payload(), template_path)

    calculator = YieldCalculator()
    with patch.object(calculator, "_calculate_is_moles", return_value=1e-4):
        config = calculator.parse_yield_config(template_path)

    assert config.is_smiles == _updated_gc_ms_yield()["internal_standard_smiles"]
    assert config.calc_method == "标准曲线"
    assert config.curve_slope == 1.25
    assert config.curve_intercept == 0.05
    assert len(config.products) == 2
    assert config.products[0].name == "对叔丁基苯甲醛"
    assert config.products[1].smiles == "CC1=CC=C(C=O)C=C1"


def test_reaction_template_without_gc_sheet_returns_default_gc_config(tmp_path: Path) -> None:
    """
    功能:
        验证旧版模板缺少 GC产率计算 Sheet 时仍可读取, 并返回默认 GC 配置.
    """
    template_path = tmp_path / "legacy_reaction_template.xlsx"
    _create_legacy_template_without_gc_sheet(template_path)

    result = read_reaction_template(template_path)

    assert result["params"]["实验名称"] == "旧版任务"
    assert result["params"]["实验ID"] == 17
    assert result["has_gc_ms_yield_sheet"] is False
    assert result["gc_ms_yield"]["internal_standard_smiles"] == ""
    assert result["gc_ms_yield"]["yield_method"] == "ECN"
    assert result["gc_ms_yield"]["products"] == []


def test_batch_in_template_read_write_round_trip(tmp_path: Path) -> None:
    """
    功能:
        验证上料文件读取, Web 数据写回和再次读取保持一致.
    """
    batch_in_path = tmp_path / "batch_in_tray.xlsx"
    _create_batch_in_template(batch_in_path)

    original = read_batch_in_template(batch_in_path)
    assert original["headers"] == ["position", "tray_type", "content", "shelf_position", "storage"]
    assert original["tray_type_options"] == [
        "2 mL试剂瓶托盘(201000705) [A1-F8]",
        "30 mL粉桶托盘(201000710) [A1-B1]",
    ]
    assert original["rows"] == [
        ["TB-2-1", "2 mL试剂瓶托盘(201000705)", "A1|乙腈|1mL", "3-1", "乙腈|A柜"]
    ]

    saved = write_batch_in_template(_updated_batch_in_payload(), batch_in_path)
    assert saved["rows"] == _updated_batch_in_payload()["rows"]
    assert saved["tray_type_options"] == original["tray_type_options"]

    workbook = openpyxl.load_workbook(batch_in_path, data_only=False)
    try:
        worksheet = workbook["batch_in_tray"]
        assert worksheet["A1"].value == "position"
        assert worksheet["C1"].value == "content"
        assert worksheet["A2"].value == "TB-2-2"
        assert worksheet["C3"].value == "12"
    finally:
        workbook.close()
