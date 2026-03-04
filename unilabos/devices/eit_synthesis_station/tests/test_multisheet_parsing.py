#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
功能:
    验证合成站多工作表解析与写入逻辑.
参数:
    无.
返回:
    无.
"""

import importlib
import json
import logging
import sys
import types
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from uuid import uuid4

import openpyxl


def _load_synthesis_modules():
    """
    功能:
        以最小依赖方式加载 synthesis manager/controller 相关模块.
        避免执行 manager 包内与外部硬件相关的初始化逻辑.
    参数:
        无.
    返回:
        Tuple[module, module, module], station_manager/task_ai_manager/station_controller 模块.
    """
    import eit_synthesis_station

    manager_pkg = types.ModuleType("eit_synthesis_station.manager")
    manager_pkg.__path__ = [str((Path.cwd() / "eit_synthesis_station" / "manager").resolve())]
    sys.modules["eit_synthesis_station.manager"] = manager_pkg
    setattr(eit_synthesis_station, "manager", manager_pkg)

    sync_module = types.ModuleType("eit_synthesis_station.manager.synchronizer")

    class EITSynthesisWorkstation:
        """
        功能:
            测试桩, 替代真实工作站基类.
        参数:
            无.
        返回:
            无.
        """

        def __init__(self, *args, **kwargs) -> None:
            return None

    sync_module.EITSynthesisWorkstation = EITSynthesisWorkstation
    sys.modules["eit_synthesis_station.manager.synchronizer"] = sync_module

    station_manager_mod = importlib.import_module("eit_synthesis_station.manager.station_manager")
    task_ai_mod = importlib.import_module("eit_synthesis_station.manager.task_ai_manager")
    station_controller_mod = importlib.import_module("eit_synthesis_station.controller.station_controller")
    return station_manager_mod, task_ai_mod, station_controller_mod


class TestSynthesisMultiSheetParsing(unittest.TestCase):
    """
    功能:
        覆盖 batch_in 读写与模板写入的多工作表路径.
    参数:
        无.
    返回:
        无.
    """

    @staticmethod
    def _make_tmp_dir() -> Path:
        """
        功能:
            创建测试临时目录.
        参数:
            无.
        返回:
            Path, 临时目录路径.
        """
        tmp_root = Path.cwd() / "eit_synthesis_station" / "tests" / "_tmp"
        tmp_root.mkdir(parents=True, exist_ok=True)
        case_dir = tmp_root / f"multisheet_{uuid4().hex}"
        case_dir.mkdir(parents=True, exist_ok=False)
        return case_dir

    def test_batch_in_read_and_auto_load_use_non_active_sheet(self) -> None:
        """
        功能:
            验证 batch_in_tray_by_file 和 auto_load_trays_from_agv 在 active 错页时仍读取正确sheet.
        参数:
            无.
        返回:
            无.
        """
        station_manager_mod, _, station_controller_mod = _load_synthesis_modules()
        case_dir = self._make_tmp_dir()
        batch_in_path = case_dir / "batch_in_tray.xlsx"

        workbook = openpyxl.Workbook()
        ws_wrong = workbook.active
        ws_wrong.title = "说明页"
        ws_wrong.cell(row=1, column=1, value="备注")
        ws_batch = workbook.create_sheet("batch_in_tray")
        ws_batch.append(["position", "tray_type", "content", "shelf_position", "storage"])
        ws_batch.append(["TB-1-1", "2mL托盘(551000502)", "A1|乙腈|1mL", "1-1", "乙腈|L1"])
        workbook.active = 0
        workbook.save(batch_in_path)
        workbook.close()

        manager = station_manager_mod.SynthesisStationManager.__new__(
            station_manager_mod.SynthesisStationManager
        )
        manager._logger = logging.getLogger("TestSynthesisMultiSheet.ManagerRead")
        captured_rows = {}

        def _fake_build_payload(rows):
            captured_rows["rows"] = rows
            return [{"resource_list": []}]

        manager.build_batch_in_tray_payload = _fake_build_payload
        manager.batch_in_tray = lambda payload: {"success": True, "payload_len": len(payload)}
        manager._generate_batch_in_tray_template = lambda _: None

        manager_result = manager.batch_in_tray_by_file(str(batch_in_path))
        self.assertEqual(manager_result["success"], True)
        self.assertEqual(captured_rows["rows"][0][0], "TB-1-1")

        controller = station_controller_mod.SynthesisStationController.__new__(
            station_controller_mod.SynthesisStationController
        )
        controller._logger = logging.getLogger("TestSynthesisMultiSheet")
        controller.list_device_status = lambda: []
        controller.open_close_door = lambda action: {"success": True, "action": action}

        fake_agv_instance = Mock()
        fake_agv_instance.batch_transfer_materials.return_value = True
        fake_agv_module = types.ModuleType("eit_agv.controller.agv_controller")
        fake_agv_module.AGVController = Mock(return_value=fake_agv_instance)

        with patch.dict(sys.modules, {"eit_agv.controller.agv_controller": fake_agv_module}):
            load_result = controller.auto_load_trays_from_agv(
                batch_in_file=str(batch_in_path),
                block=True,
            )

        self.assertEqual(load_result["success"], True)
        self.assertEqual(load_result["total_trays"], 1)
        fake_agv_instance.batch_transfer_materials.assert_called_once()

    def test_auto_generate_batch_file_writes_target_sheet_when_active_is_wrong(self) -> None:
        """
        功能:
            验证 auto_generate_batch_in_tray_from_resource_check 写入 batch_in_tray 目标sheet, 而非active错页.
        参数:
            无.
        返回:
            无.
        """
        station_manager_mod, _, _ = _load_synthesis_modules()
        case_dir = self._make_tmp_dir()

        sheet_dir = case_dir / "sheet"
        task_dir = case_dir / "data" / "tasks" / "999"
        sheet_dir.mkdir(parents=True, exist_ok=True)
        task_dir.mkdir(parents=True, exist_ok=True)

        resource_check_path = task_dir / "resource_check.json"
        resource_check_path.write_text(
            json.dumps({"missing": ["乙腈:1mL"]}, ensure_ascii=False),
            encoding="utf-8",
        )

        chemical_path = sheet_dir / "chemical_list.xlsx"
        chem_wb = openpyxl.Workbook()
        chem_ws = chem_wb.active
        chem_ws.append(["substance", "physical_state", "storage_location"])
        chem_ws.append(["乙腈", "liquid", "L1"])
        chem_wb.save(chemical_path)
        chem_wb.close()

        batch_path = sheet_dir / "batch_in_tray.xlsx"
        batch_wb = openpyxl.Workbook()
        wrong_sheet = batch_wb.active
        wrong_sheet.title = "说明页"
        wrong_sheet.append(["备注"])
        target_sheet = batch_wb.create_sheet("batch_in_tray")
        target_sheet.append(["position", "tray_type", "content", "shelf_position", "storage"])
        batch_wb.active = 0
        batch_wb.save(batch_path)
        batch_wb.close()

        manager = station_manager_mod.SynthesisStationManager.__new__(
            station_manager_mod.SynthesisStationManager
        )
        manager._logger = logging.getLogger("TestSynthesisMultiSheet.ManagerWrite")

        with patch.object(station_manager_mod, "MODULE_ROOT", case_dir):
            manager.auto_generate_batch_in_tray_from_resource_check(task_id=999)

        check_wb = openpyxl.load_workbook(batch_path)
        check_wrong = check_wb["说明页"]
        check_target = check_wb["batch_in_tray"]
        self.assertIsNone(check_wrong.cell(row=2, column=1).value)
        self.assertIsNotNone(check_target.cell(row=2, column=1).value)
        check_wb.close()

    def test_reaction_template_writer_uses_sheet_with_experiment_headers(self) -> None:
        """
        功能:
            验证 ReactionTemplateWriter 在多工作表模板中选择含“实验编号/试剂”表头的sheet写入.
        参数:
            无.
        返回:
            无.
        """
        _, task_ai_mod, _ = _load_synthesis_modules()
        case_dir = self._make_tmp_dir()
        template_path = case_dir / "reaction_template.xlsx"
        output_path = case_dir / "reaction_output.xlsx"

        workbook = openpyxl.Workbook()
        ws_wrong = workbook.active
        ws_wrong.title = "说明页"
        ws_wrong.cell(row=1, column=1, value="说明")

        ws_template = workbook.create_sheet("实验方案设定")
        ws_template.cell(row=1, column=3, value="实验编号")
        ws_template.cell(row=1, column=4, value="试剂")
        ws_template.cell(row=1, column=5, value="试剂量")
        ws_template.cell(row=2, column=3, value=1)

        workbook.active = 0
        workbook.save(template_path)
        workbook.close()

        plan = task_ai_mod.ReactionPlan(
            experiment_name="multi_sheet_case",
            experiment_id=1,
            scale_mmol=0.2,
            reactor_type="heat",
            time_h=1.0,
            temperature_c=25.0,
            rpm=500,
            open_lid_target_temp_c=25.0,
            wait_target_temp="否",
            weigh_error_percent=3.0,
            max_weigh_error_mg=1.0,
            fixed_addition_order="否",
            auto_add_stir_bar="是",
            dilution_solvent="",
            dilution_volume_ul=None,
            internal_standard_enabled="否",
            internal_standard_name="",
            internal_standard_amount_ul_or_mg=None,
            stir_after_is_min=5.0,
            flash_filter_enabled="否",
            flash_filter_solvent="",
            flash_filter_volume_ul=None,
            sample_volume_ul=1.0,
            flash_filter_experiment_numbers="全部",
            experiments=[
                task_ai_mod.ExperimentRow(
                    global_experiment_no=1,
                    reagents=[task_ai_mod.ReagentItem(name="乙腈", amount_text="1mL")],
                )
            ],
        )

        writer = task_ai_mod.ReactionTemplateWriter(template_path)
        writer.write(plan, output_path)

        result_wb = openpyxl.load_workbook(output_path)
        result_wrong = result_wb["说明页"]
        result_template = result_wb["实验方案设定"]
        self.assertIsNone(result_wrong.cell(row=2, column=4).value)
        self.assertEqual(result_template.cell(row=2, column=4).value, "乙腈")
        result_wb.close()


if __name__ == "__main__":
    unittest.main()
